"""A location's invoices as a spreadsheet, for the accountant or the
bookkeeping software: one row per invoice, or one per line item.

Two formats, from the same typed rows:
- Excel (.xlsx), for people. Item codes and invoice numbers are stored as
  text, so Excel keeps them as printed: opened from a CSV, "0081234" became
  81234 and a 16-digit code became 1.23457E+15, and nothing in a CSV can
  prevent that. Amounts and dates are real numbers and dates.
- CSV, for importing into accounting software, which reads text as text.
  Still written to open cleanly in Excel: a UTF-8 byte-order mark (without
  it Excel reads "JALAPEÑO" as mojibake), CRLF line endings.

In both, supplier text that a spreadsheet would run as a formula stays text
(see _FORMULA_START).

Every invoice is included, whatever its status, with the status in its own
column: an invoice still waiting for review is money that was spent all the
same, and leaving it out would make the export disagree with the paper.
"""
import csv
import io
import re
import uuid
from datetime import date
from decimal import Decimal
from enum import Enum

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import audit
from app.api.deps import get_tenant_or_404
from app.auth import current_user, get_db_for_tenant
from app.config import settings
from app.models import CanonicalSku, Distributor, Invoice, InvoiceLineItem, User

router = APIRouter(prefix="/exports", tags=["exports"])


class ExportKind(str, Enum):
    invoices = "invoices"
    line_items = "line-items"


class ExportFormat(str, Enum):
    xlsx = "xlsx"
    csv = "csv"


_STATUS = {
    "received": "Processing",
    "rendering": "Processing",
    "extracting": "Processing",
    "extracted": "Read",
    "needs_review": "Needs review",
    "confirmed": "Confirmed",
    "failed": "Couldn't read",
}
_MATCH = {"auto": "Matched", "pending": "Waiting for review", "confirmed": "Confirmed", "corrected": "Corrected"}


class _Kind(str, Enum):
    """What a column holds, which decides how each format writes it."""

    text = "text"  # anything a person or a supplier wrote
    date = "date"
    money = "money"  # to the cent
    price = "price"  # at the precision billed: $0.5432/lb
    quantity = "quantity"
    integer = "integer"


# A spreadsheet runs a cell that starts with one of these as a formula, and a
# supplier's description is text nobody here wrote ("=HYPERLINK(...)" on an
# invoice would otherwise be a live link in the accountant's workbook). Only
# text cells: a credit's -5.00 is ours, and must stay a number.
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")

_INVOICE_COLUMNS = [
    ("Location", _Kind.text), ("Invoice #", _Kind.text), ("Invoice date", _Kind.date),
    ("Delivery date", _Kind.date), ("Distributor", _Kind.text), ("Subtotal", _Kind.money),
    ("Tax", _Kind.money), ("Total", _Kind.money), ("Status", _Kind.text), ("Source", _Kind.text),
    ("Received", _Kind.date), ("Link", _Kind.text),
]  # fmt: skip
_LINE_COLUMNS = [
    ("Location", _Kind.text), ("Invoice #", _Kind.text), ("Invoice date", _Kind.date),
    ("Distributor", _Kind.text), ("Line", _Kind.integer), ("Item code", _Kind.text),
    ("Description", _Kind.text), ("Pack size", _Kind.text), ("Quantity", _Kind.quantity),
    ("Unit", _Kind.text), ("Unit price", _Kind.price), ("Extended price", _Kind.money),
    ("Product", _Kind.text), ("Category", _Kind.text), ("Quantity (base unit)", _Kind.quantity),
    ("Base unit", _Kind.text), ("Price per base unit", _Kind.price), ("Match", _Kind.text),
    ("Invoice status", _Kind.text), ("Link", _Kind.text),
]  # fmt: skip


# --- CSV ----------------------------------------------------------------------


def _csv_value(kind: _Kind, value: object) -> object:
    if value is None or value == "":
        return ""
    if kind == _Kind.text:
        text = str(value)
        return "'" + text if text.startswith(_FORMULA_START) else text
    if kind == _Kind.date:
        return value.isoformat()
    if kind == _Kind.money:
        return str(value.quantize(Decimal("0.01")))
    if kind == _Kind.price:
        cents = value.quantize(Decimal("0.01"))
        return str(cents) if cents == value else format(value.normalize(), "f")
    if kind == _Kind.quantity:
        return format(value.normalize(), "f")
    return value


def _csv(columns: list[tuple[str, _Kind]], rows: list[list[object]]) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow([name for name, _ in columns])
    for row in rows:
        writer.writerow([_csv_value(kind, value) for (_, kind), value in zip(columns, row)])
    return ("﻿" + out.getvalue()).encode("utf-8")


# --- Excel --------------------------------------------------------------------

_NUMBER_FORMAT = {
    _Kind.date: "yyyy-mm-dd",
    _Kind.money: "#,##0.00",
    _Kind.price: "#,##0.00##",
    _Kind.quantity: "#,##0.####",
    _Kind.integer: "0",
}
_WIDTH = {_Kind.text: 18, _Kind.date: 12, _Kind.money: 12, _Kind.price: 12, _Kind.quantity: 10, _Kind.integer: 6}
_WIDE = {"Location": 24, "Description": 44, "Product": 32, "Link": 40}


def _xlsx_cell(sheet, kind: _Kind, value: object) -> WriteOnlyCell:
    if value is None or value == "":
        return WriteOnlyCell(sheet, value=None)
    if kind == _Kind.text:
        # Characters a workbook can't hold (OCR sometimes produces control
        # characters) are dropped rather than failing the whole export.
        cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", str(value)))
        # Always text, never a formula, however it starts; and "0081234"
        # stays "0081234".
        cell.data_type = "s"
        return cell
    cell = WriteOnlyCell(sheet, value=value)
    cell.number_format = _NUMBER_FORMAT[kind]
    return cell


def _xlsx(title: str, columns: list[tuple[str, _Kind]], rows: list[list[object]]) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title)
    for i, (name, kind) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = _WIDE.get(name, _WIDTH[kind])
    sheet.freeze_panes = "A2"
    bold = Font(bold=True)
    header = [WriteOnlyCell(sheet, value=name) for name, _ in columns]
    for cell in header:
        cell.font = bold
    sheet.append(header)
    for row in rows:
        sheet.append([_xlsx_cell(sheet, kind, value) for (_, kind), value in zip(columns, row)])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


# --- The endpoint ---------------------------------------------------------------


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "location"


def _invoice_link(invoice: Invoice) -> str:
    return f"{settings.public_base_url}/invoices/{invoice.id}?location={invoice.tenant_id}"


def _status(invoice: Invoice) -> str:
    return _STATUS.get(invoice.status.value, invoice.status.value)


@router.get("/{kind}")
def export_invoices(
    kind: ExportKind,
    tenant_id: uuid.UUID,
    start: date | None = None,
    end: date | None = None,
    distributor_id: uuid.UUID | None = None,
    format: ExportFormat = ExportFormat.csv,
    db: Session = Depends(get_db_for_tenant),
    user: User = Depends(current_user),
) -> Response:
    """start/end are inclusive, on the invoice date. With either given, an
    invoice whose date couldn't be read is left out: there's no telling
    whether it belongs in the period."""
    tenant = get_tenant_or_404(db, tenant_id)
    if start and end and start > end:
        raise HTTPException(status_code=422, detail="The start date is after the end date.")

    conditions = [Invoice.tenant_id == tenant_id]
    if start:
        conditions.append(Invoice.invoice_date >= start)
    if end:
        conditions.append(Invoice.invoice_date <= end)
    if distributor_id:
        conditions.append(Invoice.distributor_id == distributor_id)
    order = (Invoice.invoice_date.asc().nulls_last(), Invoice.created_at.asc())

    if kind == ExportKind.invoices:
        columns = _INVOICE_COLUMNS
        found = db.execute(
            select(Invoice, Distributor.name)
            .outerjoin(Distributor, Distributor.id == Invoice.distributor_id)
            .where(*conditions)
            .order_by(*order)
        ).all()
        rows = [
            [
                tenant.name,
                inv.invoice_number,
                inv.invoice_date,
                inv.delivery_date,
                distributor,
                inv.subtotal,
                inv.tax,
                inv.total,
                _status(inv),
                inv.source.value,
                inv.created_at.date(),
                _invoice_link(inv),
            ]
            for inv, distributor in found
        ]
    else:
        columns = _LINE_COLUMNS
        found = db.execute(
            select(InvoiceLineItem, Invoice, Distributor.name, CanonicalSku.name, CanonicalSku.category)
            .join(Invoice, Invoice.id == InvoiceLineItem.invoice_id)
            .outerjoin(Distributor, Distributor.id == Invoice.distributor_id)
            .outerjoin(CanonicalSku, CanonicalSku.id == InvoiceLineItem.canonical_sku_id)
            .where(*conditions)
            .order_by(*order, Invoice.id, InvoiceLineItem.line_number)
        ).all()
        rows = [
            [
                tenant.name,
                inv.invoice_number,
                inv.invoice_date,
                distributor,
                line.line_number,
                line.raw_sku,
                line.raw_description,
                line.raw_pack_size,
                line.quantity,
                line.uom,
                line.unit_price,
                line.extended_price,
                product,
                category,
                line.normalized_qty_base,
                line.base_uom.value if line.base_uom else None,
                line.normalized_unit_price,
                _MATCH.get(line.review_status.value, line.review_status.value),
                _status(inv),
                _invoice_link(inv),
            ]
            for line, inv, distributor, product, category in found
        ]

    audit.record(
        db,
        user,
        "invoice.exported",
        "tenant",
        tenant_id,
        tenant_id,
        kind=kind.value,
        format=format.value,
        start=start,
        end=end,
        distributor_id=distributor_id,
        rows=len(rows),
    )
    db.commit()

    period = f"{start or 'start'}-to-{end or 'now'}" if (start or end) else "all"
    filename = f"{kind.value}-{_slug(tenant.name)}-{period}.{format.value}"
    if format == ExportFormat.xlsx:
        body = _xlsx("Invoices" if kind == ExportKind.invoices else "Line items", columns, rows)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        body = _csv(columns, rows)
        media_type = "text/csv; charset=utf-8"
    return Response(
        body,
        media_type=media_type,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            # Invoices and prices: not for a shared machine's browser cache.
            "Cache-Control": "no-store",
        },
    )
